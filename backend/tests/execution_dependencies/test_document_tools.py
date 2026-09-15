import asyncio
import io
import json
import zipfile
from uuid import uuid4

import psutil
import pytest
from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.execution_dependencies import document_tools as module
from app.execution_dependencies.attachment_tools import AttachmentBlob
from app.execution_dependencies.document_tools import READ_DOCUMENT_DEFINITION, DocumentParser, document_tool_binding
from app.infrastructure.errors import AccessDenied
from app.modules.tool.public import CallScope, ResolvedTool, ToolCall, ToolDefinition


def pdf():
    stream = b"BT /F1 12 Tf 72 720 Td (pdf document marker) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    start = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    return bytes(output)


def document(kind):
    buffer = io.BytesIO()
    if kind == "pdf":
        return pdf()
    if kind == "text":
        return b"text document marker"
    if kind == "docx":
        value = Document()
        value.add_paragraph("docx document marker")
        value.add_table(rows=1, cols=2).cell(0,0).text = "table cell"
        value.sections[0].header.paragraphs[0].text = "header marker"
        value.sections[0].footer.paragraphs[0].text = "footer marker"
        box, paragraph, run, text = (OxmlElement(name) for name in ("w:txbxContent","w:p","w:r","w:t"))
        text.text = "textbox marker"
        run.append(text)
        paragraph.append(run)
        box.append(paragraph)
        value.element.body.append(box)
        value.save(buffer)
    elif kind == "xlsx":
        value = Workbook()
        value.active.append(["xlsx document marker", 0, False])
        value.save(buffer)
        value.close()
    else:
        value = Presentation()
        slide = value.slides.add_slide(value.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "pptx document marker"
        value.save(buffer)
    return buffer.getvalue()


def context(blob, parser):
    scope = CallScope(uuid4(), uuid4(), uuid4())
    async def reader(actual, *, reference):
        assert actual == scope and reference == "files/document"
        return blob
    binding = document_tool_binding(reader, parser=parser)
    tool = ResolvedTool(ToolDefinition(uuid4(), scope.tenant_id, READ_DOCUMENT_DEFINITION), None)
    return scope, tool, binding.executor


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "pptx", "text"])
async def test_real_document_formats_are_parsed_in_terminated_workers(kind, monkeypatch):
    parser = DocumentParser()
    pids = []
    original = parser._monitor
    async def monitor(process):
        pids.append(process.pid)
        return await original(process)
    monkeypatch.setattr(parser, "_monitor", monitor)
    scope, tool, executor = context(AttachmentBlob("document." + ("txt" if kind == "text" else kind),
        "application/octet-stream", document(kind)), parser)
    try:
        result = await executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "success", result.content_json
        data = json.loads(result.content_json)
        assert kind + " document marker" in data["text"] and data["format"] == kind
        assert data["ocr"] is False and data["truncated"] is False
        if kind == "xlsx":
            assert "\t0\tFalse" in data["text"]
        if kind == "docx":
            assert all(marker in data["text"] for marker in ("table cell","header marker","footer marker","textbox marker"))
        assert not parser.active_pids and pids and all(not psutil.pid_exists(pid) for pid in pids)
    finally:
        await parser.close()


async def test_pages_and_explicit_extraction_limit_do_not_claim_full_document():
    parser = DocumentParser()
    text = "中" * 300000
    scope, tool, executor = context(AttachmentBlob("large.txt", "text/plain", text.encode()), parser)
    try:
        first = await executor.execute(tool, ToolCall("first", "read_document", '{"reference":"files/document"}'), scope)
        value = json.loads(first.content_json)
        assert value["truncated"] and value["reason"] == "text_limit" and value["next_offset"] == 16000
        last = await executor.execute(tool, ToolCall("last", "read_document", json.dumps({"reference":"files/document", "content_offset":256000})), scope)
        value = json.loads(last.content_json)
        assert value["next_offset"] is None and value["truncated"] and len(value["text"]) == 6144
    finally:
        await parser.close()


@pytest.mark.parametrize("member_size,count", [(9*1024*1024,1),(7*1024*1024,5),(0,2049)])
async def test_zip_expansion_and_member_counts_are_rejected_before_vendor_parse(member_size, count):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index in range(count):
            archive.writestr(f"part{index}.xml", b"x" * member_size)
    parser = DocumentParser()
    scope, tool, executor = context(AttachmentBlob("malicious.docx", "application/octet-stream", buffer.getvalue()), parser)
    try:
        result = await executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "error" and "archive_limit" in result.content_json
        assert not parser.active_pids
    finally:
        await parser.close()


async def test_denied_source_never_spawns_parser():
    parser = DocumentParser()
    scope = CallScope(uuid4(), uuid4(), uuid4())
    async def reader(actual, *, reference):
        raise AccessDenied("Document source is not authorized")
    binding = document_tool_binding(reader, parser=parser)
    tool = ResolvedTool(ToolDefinition(uuid4(), scope.tenant_id, READ_DOCUMENT_DEFINITION), None)
    try:
        result = await binding.executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "error" and not parser.active_pids
    finally:
        await parser.close()


@pytest.mark.parametrize("mode", ["timeout", "cancel", "close"])
async def test_hung_parser_timeout_cancel_and_close_wait_for_actual_process_exit(tmp_path, monkeypatch, mode):
    worker = tmp_path / "hung_worker.py"
    worker.write_text("import sys,time\nsys.stdin.buffer.read()\ntime.sleep(30)\n")
    monkeypatch.setattr(module, "WORKER", worker)
    parser = DocumentParser(timeout_seconds=.2 if mode == "timeout" else 15)
    scope, tool, executor = context(AttachmentBlob("file.txt", "text/plain", b"text"), parser)
    task = asyncio.create_task(executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope))
    pids = ()
    try:
        async with asyncio.timeout(3):
            while not parser.active_pids:
                await asyncio.sleep(.01)
        pids = parser.active_pids
        if mode == "cancel":
            task.cancel()
            await asyncio.sleep(0)
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task
        else:
            if mode == "close":
                await parser.close()
            result = await task
            assert result.status == "error"
        assert not parser.active_pids and all(not psutil.pid_exists(pid) for pid in pids)
    finally:
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
        await parser.close()


async def test_parser_admission_is_bounded_while_two_processes_are_busy(tmp_path, monkeypatch):
    worker = tmp_path / "hung_worker.py"
    worker.write_text("import sys,time\nsys.stdin.buffer.read()\ntime.sleep(30)\n")
    monkeypatch.setattr(module, "WORKER", worker)
    parser = DocumentParser(max_waiting=0)
    scope, tool, executor = context(AttachmentBlob("file.txt", "text/plain", b"text"), parser)
    tasks = [asyncio.create_task(executor.execute(tool, ToolCall(str(index), "read_document", '{"reference":"files/document"}'), scope)) for index in range(2)]
    try:
        async with asyncio.timeout(3):
            while len(parser.active_pids) != 2:
                await asyncio.sleep(.01)
        third = await executor.execute(tool, ToolCall("third", "read_document", '{"reference":"files/document"}'), scope)
        assert third.status == "error" and "parser_busy" in third.content_json
    finally:
        for task in tasks:
            task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)
        await parser.close()


async def test_non_bmp_text_page_fits_the_complete_worker_protocol():
    parser = DocumentParser()
    scope, tool, executor = context(AttachmentBlob("emoji.txt", "text/plain", ("🙂" * 16001).encode()), parser)
    try:
        result = await executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "success", result.content_json
        assert json.loads(result.content_json)["next_offset"] == 16000
    finally:
        await parser.close()


async def test_oversized_worker_stdout_is_killed_before_unbounded_materialization(tmp_path, monkeypatch):
    worker = tmp_path / "large_output.py"
    worker.write_text("import sys,time\nsys.stdin.buffer.read()\nsys.stdout.buffer.write(b'x'*(4*1024*1024))\nsys.stdout.flush()\ntime.sleep(30)\n")
    monkeypatch.setattr(module, "WORKER", worker)
    parser = DocumentParser()
    scope, tool, executor = context(AttachmentBlob("file.txt", "text/plain", b"data"), parser)
    try:
        result = await executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "error" and json.loads(result.content_json)["code"] == "output_limit"
        assert not parser.active_pids
    finally:
        await parser.close()


async def test_resource_supervisor_kills_real_worker_on_excess_rss_sample(tmp_path, monkeypatch):
    worker = tmp_path / "waiting.py"
    worker.write_text("import sys,time\nsys.stdin.buffer.read()\ntime.sleep(30)\n")
    monkeypatch.setattr(module, "WORKER", worker)
    actual_process = psutil.Process
    observed = []
    class ExcessRSS:
        def __init__(self, pid):
            self.process = actual_process(pid)
            observed.append(pid)
        def memory_info(self):
            sample = self.process.memory_info()
            return sample._replace(rss=512*1024*1024+1)
    monkeypatch.setattr(module.psutil, "Process", ExcessRSS)
    parser = DocumentParser()
    scope, tool, executor = context(AttachmentBlob("file.txt", "text/plain", b"data"), parser)
    try:
        result = await executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "error" and json.loads(result.content_json)["code"] == "resource_limit"
        assert observed and all(not psutil.pid_exists(pid) for pid in observed)
    finally:
        await parser.close()


@pytest.mark.parametrize("mode", ["timeout", "cancel"])
async def test_saturated_stdout_is_drained_after_kill_and_parser_slot_is_reusable(tmp_path, monkeypatch, mode):
    started, finished = tmp_path / "started", tmp_path / "finished"
    worker = tmp_path / "saturated.py"
    worker.write_text("import os,time\nfrom pathlib import Path\n"
        f"Path({str(started)!r}).touch()\nos.write(1,b'x'*(4*1024*1024))\nPath({str(finished)!r}).touch()\ntime.sleep(30)\n")
    original_worker = module.WORKER
    monkeypatch.setattr(module, "WORKER", worker)
    parser = DocumentParser(max_parallel=1, max_waiting=0, timeout_seconds=.8 if mode == "timeout" else 15)
    original_exchange = parser._exchange
    processes = []
    async def stalled_exchange(process, content, kind, offset):
        processes.append(process)
        await asyncio.Event().wait()
        return b""
    monkeypatch.setattr(parser, "_exchange", stalled_exchange)
    scope, tool, executor = context(AttachmentBlob("file.txt", "text/plain", b"reusable"), parser)
    task = asyncio.create_task(executor.execute(tool, ToolCall("read", "read_document", '{"reference":"files/document"}'), scope))
    try:
        async with asyncio.timeout(3):
            while not started.exists():
                await asyncio.sleep(.01)
        await asyncio.sleep(.05)
        assert not finished.exists()
        if mode == "cancel":
            task.cancel()
            await asyncio.sleep(0)
            task.cancel()
        done, _ = await asyncio.wait((task,), timeout=3)
        assert task in done, "Killed worker remains blocked on a saturated stdout pipe"
        if mode == "cancel":
            with pytest.raises(asyncio.CancelledError):
                await task
        else:
            result = await task
            assert json.loads(result.content_json)["code"] == "parser_timeout"
        assert processes and all(not psutil.pid_exists(process.pid) for process in processes)
        assert not parser.active_pids
        monkeypatch.setattr(module, "WORKER", original_worker)
        monkeypatch.setattr(parser, "_exchange", original_exchange)
        result = await executor.execute(tool, ToolCall("again", "read_document", '{"reference":"files/document"}'), scope)
        assert result.status == "success", result.content_json
    finally:
        task.cancel()
        if not task.done() and processes:
            # Release a failed implementation's pipe so the regression itself does not leak a worker.
            process = processes[0]
            if process.returncode is None:
                process.kill()
            if process.stdout is not None:
                while await process.stdout.read(65536):
                    pass
        await asyncio.gather(task, return_exceptions=True)
        await parser.close()
