"""Isolated document text extraction; stdin bytes and bounded stdout JSON only."""

import io
import json
import sys
import zipfile
from typing import cast

MAX_INPUT = 4 * 1024 * 1024
MAX_TEXT = 262144
MAX_UNITS = 100


class Rejected(Exception):
    pass


class Text:
    def __init__(self) -> None:
        self.parts: list[str] = []
        self.size = 0
        self.truncated = False
        self.reason: str | None = None
        self.units = 0

    def add(self, value: str) -> bool:
        if not value:
            return True
        piece = value + "\n"
        available = MAX_TEXT - self.size
        self.parts.append(piece[:available])
        self.size += min(len(piece), available)
        if len(piece) >= available:
            self.cut("text_limit")
            return False
        return True

    def cut(self, reason: str) -> None:
        self.truncated, self.reason = True, self.reason or reason


def check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 2048 or len({item.filename for item in members}) != len(members):
            raise Rejected("archive_limit")
        total = 0
        for item in members:
            if item.flag_bits & 1 or item.file_size > 8 * 1024 * 1024:
                raise Rejected("archive_limit")
            total += item.file_size
            if total > 32 * 1024 * 1024:
                raise Rejected("archive_limit")


def table(rows, text: Text, *, maximum_cells: int = 100000) -> bool:
    cells = 0
    for row in rows:
        line = []
        for value in row:
            cells += 1
            if cells > maximum_cells:
                text.cut("structure_limit")
                return False
            line.append("" if value is None else str(value))
        if not text.add("\t".join(line)):
            return False
    return True


def extract(data: bytes, kind: str) -> Text:
    text = Text()
    if kind == "text":
        text.add(data.decode("utf-8"))
        text.units = 1
        return text
    if kind in {"docx", "xlsx", "pptx"}:
        check_zip(data)
    if kind == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as document:
            if len(document.pages) > MAX_UNITS:
                text.cut("page_limit")
            for number, page in enumerate(document.pages[:MAX_UNITS], 1):
                text.units += 1
                if not text.add(f"[Page {number}]") or not text.add(page.extract_text() or ""):
                    break
                for values in page.extract_tables():
                    if not table(values, text):
                        break
                page.close()
                if text.size >= MAX_TEXT:
                    break
        return text
    if kind == "docx":
        from docx import Document
        from docx.oxml.ns import qn
        document = Document(io.BytesIO(data))
        text.units = 1
        if len(document.paragraphs) > 10000 or len(document.tables) > 100:
            text.cut("structure_limit")
        for paragraph in document.paragraphs[:10000]:
            if not text.add(paragraph.text):
                return text
        for values in document.tables[:100]:
            if not table(([cell.text for cell in row.cells] for row in values.rows), text):
                return text
        text_nodes = 0
        for shape in document.element.body.iter(qn("w:txbxContent")):
            for child in shape.iter(qn("w:t")):
                text_nodes += 1
                if text_nodes > 10000:
                    text.cut("structure_limit")
                    return text
                if not text.add(child.text or ""):
                    return text
        if len(document.sections) > 100:
            text.cut("structure_limit")
        for section in document.sections[:100]:
            for part in (section.header, section.footer):
                if part.is_linked_to_previous:
                    continue
                if len(part.paragraphs) > 10000:
                    text.cut("structure_limit")
                for paragraph in part.paragraphs[:10000]:
                    if not text.add(paragraph.text):
                        return text
        return text
    if kind == "xlsx":
        from openpyxl import load_workbook
        document = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            if len(document.sheetnames) > 32:
                text.cut("structure_limit")
            for name in document.sheetnames[:32]:
                sheet = document[name]
                text.units += 1
                if not text.add(f"[Sheet {name}]"):
                    break
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 256:
                    text.cut("structure_limit")
                if not table(sheet.iter_rows(max_row=min(sheet.max_row or 10000,10000),
                        max_col=min(sheet.max_column or 256,256), values_only=True), text):
                    break
        finally:
            document.close()
        return text
    if kind == "pptx":
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape
        from pptx.shapes.graphfrm import GraphicFrame
        document = Presentation(io.BytesIO(data))
        if len(document.slides) > MAX_UNITS:
            text.cut("page_limit")
        for index, slide in enumerate(document.slides):
            if index >= MAX_UNITS or text.size >= MAX_TEXT:
                break
            text.units += 1
            if not text.add(f"[Slide {index + 1}]"):
                break
            if len(slide.shapes) > 1000:
                text.cut("structure_limit")
            for shape_index, shape in enumerate(slide.shapes):
                if shape_index >= 1000:
                    break
                if shape.has_text_frame and not text.add(cast(Shape, shape).text_frame.text):
                    break
                if shape.has_table and not table(([cell.text for cell in row.cells] for row in cast(GraphicFrame, shape).table.rows), text):
                    break
        return text
    raise Rejected("unsupported_format")


def main() -> None:
    try:
        if sys.platform == "linux":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
        header = sys.stdin.buffer.readline(2049)
        if len(header) > 2048:
            raise Rejected("input_limit")
        request = json.loads(header)
        if not isinstance(request, dict) or set(request) != {"version", "format", "content_offset"} or request["version"] != 1:
            raise Rejected("invalid_request")
        offset = request["content_offset"]
        if type(offset) is not int or not 0 <= offset <= MAX_TEXT:
            raise Rejected("invalid_offset")
        data = sys.stdin.buffer.read(MAX_INPUT + 1)
        if len(data) > MAX_INPUT:
            raise Rejected("input_limit")
        extracted = extract(data, request["format"])
        content = "".join(extracted.parts)
        if offset > len(content):
            raise Rejected("invalid_offset")
        end = min(len(content), offset + 16000)
        result = {"version":1,"status":"success","text":content[offset:end],"content_offset":offset,
            "next_offset":end if end < len(content) else None,"offset_unit":"unicode_codepoints",
            "extracted_codepoints":len(content),"truncated":extracted.truncated,"reason":extracted.reason,
            "units_processed":extracted.units,"ocr":False}
    except Rejected as exc:
        result = {"version":1,"status":"error","code":str(exc)}
    except MemoryError:
        result = {"version":1,"status":"error","code":"resource_limit"}
    except Exception:  # noqa: BLE001 -- untrusted parser failures cross this process boundary only as a fixed code.
        result = {"version":1,"status":"error","code":"invalid_document"}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=True, separators=(",", ":")).encode("ascii"))


if __name__ == "__main__":
    main()
